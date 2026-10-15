"""ss_common.py - shared core for the "JMP + SPSS Analysis" skill (JMP/SPSS-style analysis without JMP/SPSS).

Only requires numpy, pandas, scipy, matplotlib, python-docx (+ optional: pyreadstat, python-bidi).
"""
from __future__ import annotations
import json, os, re, sys, warnings
from dataclasses import dataclass, field
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

warnings.filterwarnings("ignore")

# ----------------------------------------------------------------------------
# Language (EN / HE).  English text is the key; Hebrew is looked up if lang == "he".
# ----------------------------------------------------------------------------
HE = {
    "Mean": "ממוצע", "Std Dev": "סטיית תקן", "Std Err Mean": "שגיאת תקן של הממוצע",
    "Median": "חציון", "Minimum": "מינימום", "Maximum": "מקסימום", "N": "N",
    "Variance": "שונות", "Skewness": "אסימטריה", "Kurtosis": "גבנוניות (עודף)",
    "Quantiles": "אחוזונים", "Moments": "מומנטים", "Summary Statistics": "סטטיסטיקה תיאורית",
    "Normality tests": "בדיקות נורמליות", "Distribution": "התפלגות", "Histogram": "היסטוגרמה",
    "Normal Q-Q plot": "גרף Q-Q נורמלי", "Box plot": "תרשים קופסה", "Count": "ספירה",
    "Frequency": "שכיחות", "Value": "ערך", "Statistic": "סטטיסטי", "p-value": "ערך p",
    "Source": "מקור", "DF": "דרגות חופש", "Sum of Squares": "סכום ריבועים",
    "Mean Square": "ממוצע ריבועים", "F Ratio": "יחס F", "Prob > F": "Prob > F",
    "Model": "מודל", "Error": "שגיאה", "Total": "סה\"כ", "Residual": "שארית",
    "Term": "איבר", "Estimate": "אומדן", "Std Error": "שגיאת תקן", "t Ratio": "יחס t",
    "Prob > |t|": "Prob > |t|", "Lower 95%": "גבול תחתון 95%", "Upper 95%": "גבול עליון 95%",
    "Intercept": "קבוע", "R-squared": "R בריבוע", "Adj R-squared": "R בריבוע מתוקן",
    "RMSE": "RMSE", "Observations": "תצפיות", "Summary of Fit": "סיכום התאמה",
    "Analysis of Variance": "ניתוח שונות", "Parameter Estimates": "אומדני פרמטרים",
    "Effect Tests": "בדיקות אפקטים", "Residuals": "שאריות", "Fitted": "ערכים מוערכים",
    "Actual": "ערך בפועל", "Predicted": "ערך חזוי", "Actual by Predicted": "בפועל מול חזוי",
    "Residual by Predicted": "שארית מול חזוי", "Group": "קבוצה", "Level": "רמה",
    "Means and Std Deviations": "ממוצעים וסטיות תקן", "Tukey HSD comparisons": "השוואות Tukey HSD",
    "Pair": "זוג", "Difference": "הפרש", "Levene test": "מבחן Levene", "Welch ANOVA": "ANOVA של Welch",
    "Kruskal-Wallis": "קרוסקל-וואליס", "Oneway Analysis": "ניתוח חד-כיווני", "by": "לפי",
    "Correlation matrix": "מטריצת מתאמים", "Scatterplot Matrix": "מטריצת פיזור",
    "Principal Components": "רכיבים עיקריים", "Eigenvalue": "ערך עצמי",
    "Percent": "אחוז", "Cum Percent": "אחוז מצטבר", "Scree plot": "גרף Scree", "Loadings": "מטענים",
    "Score plot": "גרף ציונים", "Clusters": "אשכולות", "Cluster": "אשכול",
    "Cluster sizes": "גודל אשכולות", "Dendrogram": "דנדרוגרמה",
    "Control Chart": "תרשים בקרה", "UCL": "UCL", "LCL": "LCL", "Center line": "קו מרכזי",
    "Subgroup": "תת-קבוצה", "Individual": "ערך בודד", "Moving Range": "טווח נע", "Range": "טווח",
    "Out-of-control points": "נקודות מחוץ לבקרה", "Process Capability": "יכולת תהליך",
    "LSL": "גבול מפרט תחתון (LSL)", "USL": "גבול מפרט עליון (USL)", "Target": "יעד",
    "Within sigma": "סיגמה בתוך תת-קבוצות", "Overall sigma": "סיגמה כוללת",
    "PPM below LSL": "PPM מתחת ל-LSL", "PPM above USL": "PPM מעל USL", "PPM total": "PPM כולל",
    "Gage R&R": "Gage R&R", "Repeatability": "חזרתיות", "Reproducibility": "שחזוריות",
    "Part-to-Part": "בין חלקים", "Operator": "מפעיל", "Part": "חלק", "Measurement": "מדידה",
    "Variance Component": "רכיב שונות", "% of Total Variance": "% מסך השונות",
    "StdDev": "סטיית תקן", "% Study Var": "% שונות מחקר", "% Tolerance": "% סובלנות",
    "Number of Distinct Categories": "מספר קטגוריות נבדלות", "Variability Chart": "תרשים שונות",
    "Verdict": "מסקנה", "Weibull": "וייבול", "Lognormal": "לוג-נורמלי", "Exponential": "מעריכי",
    "Shape (beta)": "פרמטר צורה (β)", "Scale (eta)": "פרמטר קנה מידה (η)", "Log-Likelihood": "לוג-נראות",
    "AICc": "AICc", "BIC": "BIC", "Failures": "כשלים", "Censored": "מצונזרים", "Time": "זמן",
    "Survival": "הישרדות", "Kaplan-Meier": "קפלן-מאייר", "Log-rank test": "מבחן Log-rank",
    "Probability plot": "גרף הסתברות", "Cumulative failure": "כשל מצטבר", "B10 life": "חיים B10",
    "Mean life (MTTF)": "זמן חיים ממוצע (MTTF)", "Reliability at time": "אמינות בזמן",
    "Main Effects": "אפקטים ראשיים", "Interaction": "אינטראקציה", "Effect": "אפקט",
    "Coefficient": "מקדם", "Pareto of Effects": "פרטו של אפקטים", "Half-normal plot": "גרף חצי-נורמלי",
    "Design matrix": "מטריצת תכנון", "Run": "ריצה", "Factors": "גורמים", "Response": "תגובה",
    "Interaction plot": "גרף אינטראקציה", "Lenth PSE": "Lenth PSE", "Sample size": "גודל מדגם",
    "t-test": "מבחן t", "Paired t-test": "מבחן t מזווג", "Two-sample t-test": "מבחן t לשני מדגמים",
    "One-sample t-test": "מבחן t למדגם אחד", "Mann-Whitney U": "מאן-וויטני U",
    "Wilcoxon signed-rank": "ווילקוקסון", "Chi-square test": "מבחן חי בריבוע",
    "Contingency table": "טבלת שכיחויות", "Cramer's V": "V של קרמר", "Expected": "צפוי", "Observed": "נצפה",
    "Notes": "הערות", "Data": "נתונים", "Report": "דוח", "Generated by": "הופק על ידי",
    "Logistic Regression": "רגרסיה לוגיסטית", "Odds Ratio": "יחס סיכויים", "Coefficient (log-odds)": "מקדם (log-odds)",
    "Accuracy": "דיוק", "ROC curve": "עקומת ROC", "AUC": "AUC", "Confusion matrix": "מטריצת בלבול",
    "Interval": "רווח סמך", "Mean Diff": "הפרש ממוצעים", "Conclusion": "מסקנה",
    "Significant at 0.05": "מובהק ברמת 0.05", "Not significant at 0.05": "לא מובהק ברמת 0.05",
    "Method": "שיטה", "Order": "סדר", "Variable": "משתנה", "Type III SS": "SS מסוג III",
    "Studentized Residuals": "שאריות מתוקננות", "Leverage": "מינוף", "Cook's D": "מרחק קוק",
    "Residual normal Q-Q": "Q-Q של שאריות", "Residual histogram": "היסטוגרמת שאריות",
    "Observation order": "סדר תצפית", "Effect size": "גודל אפקט", "Cohen's d": "d של כהן",
    "Power": "עוצמה", "Alpha": "אלפה", "Normal": "נורמלי", "Shapiro-Wilk": "שאפירו-וילק",
    "Anderson-Darling": "אנדרסון-דרלינג", "Lilliefors / KS": "קולמוגורוב-סמירנוב",
    "Sample": "מדגם", "Capability histogram": "היסטוגרמת יכולת",
}

PALETTE = ["#1F5AA6", "#D9822B", "#3B9B6E", "#B03A48", "#7B5EA7", "#8C6D46", "#2B8CA3", "#7F7F7F"]


class Lang:
    def __init__(self, lang: str = "en"):
        self.lang = "he" if str(lang).lower().startswith("he") else "en"

    @property
    def rtl(self):
        return self.lang == "he"

    def t(self, s: str) -> str:
        return HE.get(s, s) if self.rtl else s

    def fig(self, s: str) -> str:
        """text for matplotlib (visual order for Hebrew)."""
        return visual(self.t(s)) if self.rtl else s


_HEB = re.compile(r"[֐-׿]")


def visual(text: str) -> str:
    """Convert logical-order Hebrew text to visual order for matplotlib."""
    if not _HEB.search(text):
        return text
    try:
        from bidi.algorithm import get_display
        return get_display(text)
    except Exception:
        pass
    words = text.split(" ")
    out = []
    for w in words:
        out.append(w[::-1] if _HEB.search(w) else w)
    return " ".join(out[::-1])


# ----------------------------------------------------------------------------
# Plot style (JMP-like: clean, light grid, blue markers)
# ----------------------------------------------------------------------------
def set_style():
    plt.rcParams.update({
        "figure.dpi": 110, "savefig.dpi": 160, "font.family": "DejaVu Sans",
        "axes.spines.top": False, "axes.spines.right": False, "axes.grid": True,
        "grid.color": "#DDDDDD", "grid.linewidth": 0.6, "axes.edgecolor": "#555555",
        "axes.titleweight": "bold", "axes.titlesize": 11, "axes.labelsize": 10,
        "axes.prop_cycle": matplotlib.cycler(color=PALETTE), "legend.frameon": False,
        "figure.autolayout": False,
    })


set_style()


def save_fig(fig, outdir: str, name: str) -> str:
    os.makedirs(outdir, exist_ok=True)
    path = os.path.join(outdir, re.sub(r"[^\w\-]+", "_", name) + ".png")
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


# ----------------------------------------------------------------------------
# Data loading + profiling (JMP-style modeling types)
# ----------------------------------------------------------------------------
def load_data(path: str, sheet=None) -> pd.DataFrame:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".txt", ".tsv"):
        for enc in ("utf-8-sig", "cp1255", "latin-1"):
            try:
                with open(path, encoding=enc) as fh:
                    head = fh.readline()
                sep = "\t" if "\t" in head else (";" if head.count(";") > head.count(",") else ",")
                return pd.read_csv(path, sep=sep, encoding=enc)
            except UnicodeDecodeError:
                continue
        raise ValueError("Could not read text file")
    if ext in (".xlsx", ".xlsm", ".xls"):
        return pd.read_excel(path, sheet_name=sheet or 0)
    if ext == ".sav":
        try:
            import pyreadstat
            df, _ = pyreadstat.read_sav(path)
            return df
        except ImportError:
            return pd.read_spss(path)  # needs pyreadstat as well
    if ext == ".json":
        return pd.read_json(path)
    if ext == ".jmp":
        raise ValueError("JMP (.jmp) files cannot be read directly. In JMP: File > Save As > Excel/CSV, then pass that file.")
    raise ValueError(f"Unsupported file type: {ext}")


def profile(df: pd.DataFrame, max_levels: int = 12) -> pd.DataFrame:
    """Modeling-type guess per column: continuous / ordinal / nominal."""
    rows = []
    for c in df.columns:
        s = df[c]
        nun = s.nunique(dropna=True)
        if pd.api.types.is_numeric_dtype(s):
            mt = "ordinal/discrete" if nun <= max_levels and (s.dropna() % 1 == 0).all() else "continuous"
        else:
            mt = "nominal"
        rows.append({"column": c, "dtype": str(s.dtype), "modeling_type": mt, "n_unique": nun,
                     "n_missing": int(s.isna().sum()), "example": s.dropna().iloc[0] if s.notna().any() else None})
    return pd.DataFrame(rows)


def numeric_cols(df):
    return [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c])]


# ----------------------------------------------------------------------------
# Result container + Word report
# ----------------------------------------------------------------------------
@dataclass
class Result:
    title: str
    tables: list = field(default_factory=list)   # (name, DataFrame)
    figures: list = field(default_factory=list)  # (caption, path)
    notes: list = field(default_factory=list)
    summary: dict = field(default_factory=dict)  # machine-readable key numbers

    def table(self, name, df, index=False):
        self.tables.append((name, df.reset_index() if index else df))

    def fig(self, caption, path):
        self.figures.append((caption, path))

    def note(self, text):
        self.notes.append(text)


def fmt_p(p):
    if p is None or (isinstance(p, float) and np.isnan(p)):
        return ""
    return "<.0001" if p < 1e-4 else f"{p:.4f}"


def _fmt(v, nd=4):
    if isinstance(v, (float, np.floating)):
        if np.isnan(v):
            return ""
        if abs(v) >= 1e6 or (0 < abs(v) < 1e-4):
            return f"{v:.3e}"
        return f"{v:.{nd}g}" if abs(v) < 1000 else f"{v:,.2f}"
    return str(v)


def _rtl_par(par):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    pPr = par._p.get_or_add_pPr()
    b = OxmlElement("w:bidi")
    pPr.append(b)
    for r in par.runs:
        rPr = r._r.get_or_add_rPr()
        rt = OxmlElement("w:rtl")
        rPr.append(rt)


def build_docx(results, path, L: Lang, title=None, meta: dict | None = None):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)

    def para(text, bold=False, size=None, color=None, style=None, align=None):
        p = doc.add_paragraph(style=style)
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor(*color)
        if L.rtl:
            _rtl_par(p)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif align:
            p.alignment = align
        return p

    para(title or L.t("Report"), bold=True, size=20, color=(0x1F, 0x5A, 0xA6))
    if meta:
        for k, v in meta.items():
            para(f"{k}: {v}", size=9, color=(0x66, 0x66, 0x66))
    for res in results:
        h = doc.add_heading(res.title, level=1)
        if L.rtl:
            _rtl_par(h)
            h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for name, df in res.tables:
            para(name, bold=True, size=11)
            d = df.copy()
            tbl = doc.add_table(rows=1, cols=len(d.columns))
            tbl.style = "Light Grid Accent 1"
            cols = list(d.columns)
            if L.rtl:
                tbl.alignment = 2
            for j, c in enumerate(cols):
                cell = tbl.rows[0].cells[j]
                cell.text = L.t(str(c))
                for r in cell.paragraphs[0].runs:
                    r.bold = True
                    r.font.size = Pt(9)
            for _, row in d.iterrows():
                cells = tbl.add_row().cells
                for j, c in enumerate(cols):
                    v = row[c]
                    txt = fmt_p(v) if ("Prob" in str(c) or "p-value" in str(c) or str(c) == "p") and isinstance(v, (float, np.floating)) else _fmt(v)
                    cells[j].text = txt
                    for r in cells[j].paragraphs[0].runs:
                        r.font.size = Pt(9)
            doc.add_paragraph()
        for cap, p in res.figures:
            if os.path.exists(p):
                doc.add_picture(p, width=Cm(15.5))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                para(cap, size=9, color=(0x55, 0x55, 0x55))
        if res.notes:
            para(L.t("Notes"), bold=True, size=11)
            for n in res.notes:
                para("• " + n, size=10)
    doc.save(path)
    return path


def finish(results, outdir, L: Lang, basename="report", title=None, meta=None):
    """Write Word report, CSV tables, JSON summary. Returns dict of paths."""
    os.makedirs(outdir, exist_ok=True)
    if isinstance(results, Result):
        results = [results]
    docx_path = build_docx(results, os.path.join(outdir, basename + ".docx"), L, title or results[0].title, meta)
    summ = {}
    for i, r in enumerate(results):
        summ[r.title] = r.summary
        for name, df in r.tables:
            fn = re.sub(r"[^\w\-]+", "_", f"{basename}_{i}_{name}")[:80] + ".csv"
            df.to_csv(os.path.join(outdir, fn), index=False, encoding="utf-8-sig")
    with open(os.path.join(outdir, basename + "_summary.json"), "w", encoding="utf-8") as f:
        json.dump(summ, f, ensure_ascii=False, indent=2, default=lambda o: o.item() if hasattr(o, "item") else str(o))
    return {"docx": docx_path, "outdir": outdir,
            "figures": [p for r in results for _, p in r.figures]}


def get_param(params: dict, key, default=None, required=False):
    v = params.get(key, default)
    if required and v is None:
        raise ValueError(f"Missing required parameter: {key}")
    return v


def clean_xy(df, cols):
    d = df[cols].dropna()
    return d
